"""
가입 동아리 현황·운영진 명단(xlsx)을 읽어 기초 데이터 SQL을 만든다.

    python scripts/import-roster.py ["UNIDEV 가입 동아리 현황 및 대표자 구성.xlsx"]

만드는 파일 (실제 개인정보이므로 저장소에 올리지 않는다, .gitignore의 db/private/)
    db/private/10_roster.sql             동아리·구성원·역할 배정. 로컬과 Supabase 모두에 적용
    db/private/local/90_dev_accounts.sql 로컬 전용: 현직 운영진 전원에게 개발용 로그인 계정 연결

읽는 시트
    '2026 가입 동아리 현황…'  연번 | 등록 형태 | 학교 | 동아리 | 대표자 | 전화번호 | 이메일 | 비고
    '2026 운영진…'            연번 | 부서 | 직책 | 학교 | 동아리 | 성명 | 전화번호 | 이메일 | 비고

비고 칸에서 읽는 변경 기록
    '26.06.29 SNUGDC 대표자 교체(구윤환 -> 박세미)'
    '26.09.01 홍보팀장 교체(SATURN 김연재 -> 마나 임담희)'
    '26.09.01 행사 부팀장 임명'
    '26.09.14 동아리 공식 명칭 변경(얼간이들 -> IDIOTS)'

시작일을 알 수 없는 배정은 임기 시작일(TERM_START)로 넣고 근거 칸에 그 사실을 적는다.
"""
import re
import sys
import uuid
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'UNIDEV 가입 동아리 현황 및 대표자 구성.xlsx'
OUT = ROOT / 'db' / 'private' / '10_roster.sql'
OUT_LOCAL = ROOT / 'db' / 'private' / 'local' / '90_dev_accounts.sql'

TERM_START = '2026-03-01'
ROSTER_BASIS = '2026 운영진 명단(26.09~)'
UNKNOWN_START = f'{ROSTER_BASIS} 일괄 등록, 시작일 미상(임기 시작일로 입력)'
NS = uuid.UUID('6f1c2b8e-4d3a-4e1f-9c55-7a2d0e9b1c34')  # 같은 원본이면 같은 id가 나오도록

STATUS = {'정등록': 'full', '가등록': 'provisional', '심사 중': 'pending', '비활성': 'inactive'}
warnings: list[str] = []


def warn(msg):
    warnings.append(msg)


def s(v):
    return None if v is None or str(v).strip() == '' else str(v).strip()


def q(v):
    """SQL 문자열 리터럴"""
    if v is None:
        return 'null'
    return "'" + str(v).replace("'", "''") + "'"


def date_from(yy, mm, dd):
    return f'20{yy}-{int(mm):02d}-{int(dd):02d}'


def day_before(d):
    import datetime
    return (datetime.date.fromisoformat(d) - datetime.timedelta(days=1)).isoformat()


def find_sheet(wb, keyword):
    for ws in wb.worksheets:
        if keyword in ws.title:
            return ws
    sys.exit(f"'{keyword}'가 들어간 시트를 찾을 수 없습니다: {[w.title for w in wb.worksheets]}")


def rows_of(ws, headers):
    """첫 행을 머리글로 보고, 필요한 열만 dict로 돌려준다 (연번이 있는 행만)"""
    head = [s(c.value) for c in ws[1]]
    idx = {}
    for h in headers:
        if h not in head:
            sys.exit(f"시트 '{ws.title}'에 '{h}' 열이 없습니다. 머리글: {head}")
        idx[h] = head.index(h)
    out = []
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[idx['연번']] is None:
            continue
        rec = {h: s(row[i]) if i < len(row) else None for h, i in idx.items()}
        rec['_row'] = r
        out.append(rec)
    return out


wb = openpyxl.load_workbook(SRC, data_only=True)
club_rows = rows_of(find_sheet(wb, '가입 동아리'), ['연번', '등록 형태', '학교', '동아리', '대표자', '전화번호', '이메일', '비고'])
staff_rows = rows_of(find_sheet(wb, '운영진'), ['연번', '부서', '직책', '학교', '동아리', '성명', '전화번호', '이메일', '비고'])

# ---------------------------------------------------------------------
# 동아리
# ---------------------------------------------------------------------
clubs = {}          # 동아리명 → dict
for r in club_rows:
    name = r['동아리']
    if name in clubs:
        warn(f"가입 동아리 시트 {r['_row']}행: 동아리명 '{name}' 중복")
    status = STATUS.get(r['등록 형태'])
    if not status:
        warn(f"가입 동아리 시트 {r['_row']}행: 알 수 없는 등록 형태 '{r['등록 형태']}' → 심사 중으로 입력")
        status = 'pending'
    clubs[name] = {
        'id': str(uuid.uuid5(NS, f"club:{r['학교']}:{name}")),
        'university': r['학교'], 'name': name, 'status': status, 'note': None,
        'rep': r, 'row': r['_row'],
    }

# ---------------------------------------------------------------------
# 구성원 (이름+전화번호로 같은 사람 판단)
# ---------------------------------------------------------------------
members = {}        # key → dict


def member_key(name, phone):
    return f'{name}|{phone or ""}'


def upsert_member(name, university, club_name, phone, email, src):
    key = member_key(name, phone)
    club = clubs.get(club_name)
    if club_name and not club:
        warn(f"{src}: 동아리 '{club_name}'가 가입 동아리 시트에 없음 (소속 동아리 비워 둠)")
    if key in members:
        m = members[key]
        if email and m['email'] and email != m['email']:
            warn(f"{src}: {name} 이메일이 시트마다 다름 ({m['email']} / {email}) → {m['email']} 사용")
        m['email'] = m['email'] or email
        return m
    # 이름만 같고 전화번호가 다른 경우
    for m in members.values():
        if m['name'] == name and m['club'] == club_name and phone and m['phone'] and m['phone'] != phone:
            warn(f"{src}: {name}({club_name}) 전화번호가 시트마다 다름 ({m['phone']} / {phone}) → {m['phone']} 사용 (확인 필요)")
            return m
    m = {
        'id': str(uuid.uuid5(NS, f'member:{key}')), 'name': name, 'university': university,
        'club': club_name if club else None, 'phone': phone, 'email': email,
    }
    members[key] = m
    return m


def find_member(name, club_name):
    """비고에 이름만 나오는 사람(전임자 등). 없으면 연락처 없이 만든다."""
    for m in members.values():
        if m['name'] == name and (club_name is None or m['club'] == club_name):
            return m
    univ = clubs[club_name]['university'] if club_name in clubs else None
    m = {
        'id': str(uuid.uuid5(NS, f'member:former:{club_name}:{name}')), 'name': name, 'university': univ,
        'club': club_name if club_name in clubs else None, 'phone': None, 'email': None,
    }
    members[f'former|{club_name}|{name}'] = m
    return m


EMAIL_RE = re.compile(r'^[^@\s]+@[^@\s.]+(\.[^@\s.]+)+$')
TYPO_DOMAINS = {'gamil.com': 'gmail.com', 'gmial.com': 'gmail.com', 'naver.con': 'naver.com'}


def check_contact(name, phone, email, src):
    if email and not EMAIL_RE.match(email):
        warn(f"{src}: {name} 이메일 형식 오류 '{email}' (원본 그대로 입력, 확인 필요)")
    elif email and email.split('@')[1].lower() in TYPO_DOMAINS:
        warn(f"{src}: {name} 이메일 도메인 오타로 보임 '{email}' (원본 그대로 입력, 확인 필요)")
    if phone and not re.match(r'^01\d-\d{3,4}-\d{4}$', phone):
        warn(f"{src}: {name} 전화번호 형식 확인 필요 '{phone}'")


# 운영진 시트를 먼저 읽는다 (학교·소속 정보가 사람 기준으로 적혀 있음)
for r in staff_rows:
    src = f"운영진 시트 {r['_row']}행"
    check_contact(r['성명'], r['전화번호'], r['이메일'], src)
    club = clubs.get(r['동아리'])
    if club and r['학교'] != club['university']:
        warn(f"{src}: {r['성명']}의 학교 '{r['학교']}'가 동아리 {r['동아리']}의 학교 '{club['university']}'와 다름 "
             f"(동아리는 '{club['university']}', 구성원 소속 대학은 '{r['학교']}'로 입력, 확인 필요)")
    r['_member'] = upsert_member(r['성명'], r['학교'], r['동아리'], r['전화번호'], r['이메일'], src)

for c in clubs.values():
    r = c['rep']
    src = f"가입 동아리 시트 {r['_row']}행"
    if not r['대표자']:
        continue
    check_contact(r['대표자'], r['전화번호'], r['이메일'], src)
    c['rep_member'] = upsert_member(r['대표자'], c['university'], c['name'], r['전화번호'], r['이메일'], src)

# ---------------------------------------------------------------------
# 역할 배정
# ---------------------------------------------------------------------
assignments = []    # dict(member, role, club, team, start, end, basis, note)


def assign(member, role, *, club=None, team=None, start=TERM_START, end=None, basis=UNKNOWN_START, note=None):
    assignments.append(dict(member=member, role=role, club=club, team=team, start=start, end=end, basis=basis, note=note))


REP_CHANGE = re.compile(r'(\d{2})\.(\d{2})\.(\d{2})\.?\s+(.+?)\s*대표자 교체\((.+?)\s*->\s*(.+?)\)')
LEAD_CHANGE = re.compile(r'(\d{2})\.(\d{2})\.(\d{2})\.?\s+(\S+?)팀장 교체\((\S+)\s+(\S+)\s*->\s*(\S+)\s+(\S+)\)')
VICE_LEAD = re.compile(r'(\d{2})\.(\d{2})\.(\d{2})\.?\s+(\S+?)\s*부팀장 임명')
RENAME = re.compile(r'(\d{2})\.(\d{2})\.(\d{2})\.?\s+동아리 공식 명칭 변경\((.+?)\s*->\s*(.+?)\)')

rep_start = {}      # 동아리명 → (시작일, 근거)
handled_notes = set()

all_notes = [(f"가입 동아리 시트 {c['rep']['_row']}행", c['rep']['비고']) for c in clubs.values()] + \
            [(f"운영진 시트 {r['_row']}행", r['비고']) for r in staff_rows]
for src, note in all_notes:
    if not note or note in handled_notes:
        continue
    if m := REP_CHANGE.search(note):
        d = date_from(*m.group(1, 2, 3))
        club_name, old, new = m.group(4).strip(), m.group(5).strip(), m.group(6).strip()
        if club_name not in clubs:
            warn(f'{src}: 비고의 동아리 \'{club_name}\'를 찾을 수 없음: {note}')
            continue
        if clubs[club_name]['rep']['대표자'] != new:
            warn(f'{src}: 비고의 신임 대표자 \'{new}\'가 현재 대표자 \'{clubs[club_name]["rep"]["대표자"]}\'와 다름: {note}')
        basis = note.strip()
        rep_start[club_name] = (d, basis)
        prev = find_member(old, club_name)
        prev_start = TERM_START if TERM_START < d else d
        assign(prev, 'representative', club=club_name, start=prev_start, end=day_before(d) if day_before(d) >= prev_start else prev_start,
               basis=f'{basis} / 시작일 미상(임기 시작일로 입력)')
    elif m := LEAD_CHANGE.search(note):
        d = date_from(*m.group(1, 2, 3))
        team = m.group(4) + '팀'
        old_club, old_name, new_club, new_name = m.group(5, 6, 7, 8)
        prev = find_member(old_name, old_club)
        assign(prev, 'team_lead', team=team, start=TERM_START, end=day_before(d), basis=f'{note.strip()} / 시작일 미상(임기 시작일로 입력)')
        for r in staff_rows:
            if r['성명'] == new_name and r['동아리'] == new_club:
                r['_lead_start'] = (d, note.strip())
    elif m := VICE_LEAD.search(note):
        d = date_from(*m.group(1, 2, 3))
        for r in staff_rows:
            if r['비고'] == note:
                r['_vice_lead'] = (m.group(4).strip() + '팀', d, note.strip())
    elif m := RENAME.search(note):
        d = date_from(*m.group(1, 2, 3))
        old, new = m.group(4).strip(), m.group(5).strip()
        if new in clubs:
            clubs[new]['note'] = f'{d} 공식 명칭 변경 ({old} → {new})'
    else:
        warn(f'{src}: 해석하지 못한 비고 (반영 안 함): {note}')
        continue
    handled_notes.add(note)

# 현직 대표자 (가입 동아리 시트 기준)
for c in clubs.values():
    if 'rep_member' not in c:
        warn(f"가입 동아리 시트 {c['row']}행: {c['name']} 대표자 공란")
        continue
    start, basis = rep_start.get(c['name'], (TERM_START, UNKNOWN_START))
    assign(c['rep_member'], 'representative', club=c['name'], start=start, basis=basis)

# 운영진 시트의 직책·부서
TEAM_LEAD = re.compile(r'^(\S+팀)장$')
for r in staff_rows:
    m_ = r['_member']
    title, dept = r['직책'] or '', r['부서']
    src = f"운영진 시트 {r['_row']}행"
    lead_start, lead_basis = r.get('_lead_start', (TERM_START, UNKNOWN_START))
    if title == '회장':
        assign(m_, 'president')
    elif title.startswith('부회장'):
        assign(m_, 'vice_president')
        if (inner := re.search(r'\((\S+팀)장\)', title)):
            assign(m_, 'team_lead', team=inner.group(1), start=lead_start, basis=lead_basis)
    elif (tl := TEAM_LEAD.match(title)):
        if dept and dept != tl.group(1):
            warn(f"{src}: 직책 '{title}'과 부서 '{dept}'가 다름 → 직책 기준")
        assign(m_, 'team_lead', team=tl.group(1), start=lead_start, basis=lead_basis)
    elif title == '대표자':
        rep_of = [c for c in clubs.values() if c.get('rep_member') is m_]
        if not rep_of:
            warn(f"{src}: {r['성명']}이 대표자로 적혀 있으나 가입 동아리 시트의 대표자와 맞지 않음")
        if dept:
            if '_vice_lead' in r:
                team, d, note = r['_vice_lead']
                assign(m_, 'vice_team_lead', team=team, start=d, basis=note)
            else:
                assign(m_, 'team_member', team=dept)
    elif title == '운영진':
        if dept:
            assign(m_, 'team_member', team=dept)
        else:
            assign(m_, 'staff')
    else:
        warn(f"{src}: 알 수 없는 직책 '{title}' (반영 안 함)")

# 시트의 부서별 인원 요약표와 대조 (K~L열)
ws_staff = find_sheet(wb, '운영진')
summary = {}
for row in ws_staff.iter_rows(values_only=True):
    for i, v in enumerate(row[:-1]):
        if isinstance(v, str) and v.strip().endswith('팀') and isinstance(row[i + 1], (int, float)):
            summary[v.strip()] = int(row[i + 1])
counted = {}
for a in assignments:
    if a['team'] and a['end'] is None:
        counted.setdefault(a['team'], set()).add(a['member']['id'])
for team, n in summary.items():
    if len(counted.get(team, ())) != n:
        warn(f"부서별 인원 요약표와 불일치: {team} 요약 {n}명, 명단 {len(counted.get(team, ()))}명")

# ---------------------------------------------------------------------
# SQL 출력
# ---------------------------------------------------------------------
L = []
L.append('-- 자동 생성 파일: python scripts/import-roster.py (직접 수정하지 말고 원본 xlsx를 고친 뒤 다시 생성)')
L.append(f'-- 원본: {SRC.name}')
L.append('-- 실제 개인정보가 들어 있다. 저장소에 올리지 않는다. (db/private/는 .gitignore 대상)')
L.append('-- db/migrations, db/seed 적용 후 실행한다. (로컬: npm run db:reset / Supabase: SQL Editor)')
L.append('')
L.append('insert into public.clubs (id, university, name, status, note) values')
L.append(',\n'.join(f"  ({q(c['id'])}, {q(c['university'])}, {q(c['name'])}, {q(c['status'])}, {q(c['note'])})" for c in clubs.values()) + ';')
L.append('')
L.append('insert into public.members (id, name, university, club_id, phone, email) values')
ordered = sorted(members.values(), key=lambda m: m['name'])
L.append(',\n'.join(
    f"  ({q(m['id'])}, {q(m['name'])}, {q(m['university'])}, {q(clubs[m['club']]['id']) if m['club'] else 'null'}, {q(m['phone'])}, {q(m['email'])})"
    for m in ordered) + ';')
L.append('')
L.append('-- 역할 배정: 종료된 기록을 먼저 넣는다 (현직 중복 검사 트리거 순서)')
L.append('insert into public.role_assignments (member_id, role_id, club_id, team_id, start_on, end_on, basis)')
L.append('select a.member_id::uuid, r.id, a.club_id::uuid,')
L.append('       (select t.id from public.teams t where t.name = a.team_name and t.closed_on is null),')
L.append('       a.start_on::date, a.end_on::date, a.basis')
L.append('  from (values')
vals = []
for a in sorted(assignments, key=lambda a: (a['end'] is None, a['start'])):
    club_id = clubs[a['club']]['id'] if a['club'] else None
    vals.append(f"    ({q(a['member']['id'])}, {q(a['role'])}, {q(club_id)}, {q(a['team'])}, {q(a['start'])}, {q(a['end'])}, {q(a['basis'])})")
L.append(',\n'.join(vals))
L.append('  ) as a(member_id, role_code, club_id, team_name, start_on, end_on, basis)')
L.append('  join public.roles r on r.code = a.role_code;')
L.append('')
L.append('-- 팀 이름이 시드와 맞지 않으면 team_id가 비어 위 insert가 트리거에서 실패한다.')
OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text('\n'.join(L) + '\n', encoding='utf-8')

# 로컬 전용 개발용 로그인 계정
#   운영진 계정: 대표자·대리인 외의 현직 역할이 있는 사람. 아이디는 이메일 앞부분(없거나 겹치면 staff-N)
#   대표자 계정: 동아리마다 1개. 아이디는 동아리 영문 이름(없으면 club-N)
CLUB_ROLES = {'representative', 'proxy'}
staff_members = sorted({a['member']['id']: a['member'] for a in assignments
                        if a['end'] is None and a['role'] not in CLUB_ROLES}.values(), key=lambda m: m['name'])
used = set()


def make_username(base, fallback):
    u = re.sub(r'[^a-z0-9._-]', '', (base or '').lower()).strip('._-')
    if not re.match(r'^[a-z0-9][a-z0-9._-]{2,31}$', u) or u in used:
        u = fallback
    used.add(u)
    return u


dev_accounts = []   # (username, kind, member_id, club_id, label)
for i, m in enumerate(staff_members, start=1):
    dev_accounts.append((make_username((m['email'] or '').split('@')[0], f'staff-{i}'), 'staff', m['id'], None, m['name']))
for i, c in enumerate(clubs.values(), start=1):
    dev_accounts.append((make_username(c['name'], f'club-{i}'), 'club', None, c['id'], c['name']))
dev_accounts.append((make_username('admin', 'admin-1'), 'admin', None, None, '관리자'))

D = ['-- 로컬 전용 (Supabase에 적용하지 않음): 개발용 로그인 계정. 비밀번호는 .env의 DEV_ACCOUNT_PASSWORD',
     '-- Supabase에서는 관리자가 로그인 계정 화면에서 계정을 만든다.',
     'create temp table dev_accounts (username text, kind text, member_id uuid, club_id uuid, label text);',
     'insert into dev_accounts values']
D.append(',\n'.join(f"  ({q(u)}, {q(k)}, {q(mid)}, {q(cid)}, {q(label)})" for u, k, mid, cid, label in dev_accounts) + ';')
D += ["insert into auth.users (id, email, encrypted_password)",
      "select gen_random_uuid(), username || '@' || :'login_domain', crypt(:'dev_password', gen_salt('bf')) from dev_accounts;",
      "insert into public.accounts (auth_user_id, username, kind, member_id, club_id)",
      "select u.id, d.username, d.kind, d.member_id, d.club_id",
      "  from dev_accounts d join auth.users u on u.email = d.username || '@' || :'login_domain';"]
OUT_LOCAL.parent.mkdir(parents=True, exist_ok=True)
OUT_LOCAL.write_text('\n'.join(D) + '\n', encoding='utf-8')

# ---------------------------------------------------------------------
# 요약
# ---------------------------------------------------------------------
from collections import Counter
print(f'원본: {SRC.name}')
print(f'동아리 {len(clubs)}개 ' + str(dict(Counter(c["status"] for c in clubs.values()))))
print(f'구성원 {len(members)}명 (연락처 없는 전임자 {sum(1 for m in members.values() if not m["phone"] and not m["email"])}명 포함)')
cur = [a for a in assignments if a['end'] is None]
print(f'역할 배정 {len(assignments)}건 (현직 {len(cur)}건, 종료 {len(assignments) - len(cur)}건) ' + str(dict(Counter(a["role"] for a in cur))))
print(f'팀별 현직 인원: ' + ', '.join(f'{t} {len(v)}' for t, v in counted.items()))
print(f'생성: {OUT.relative_to(ROOT)}, {OUT_LOCAL.relative_to(ROOT)} (개발용 계정: 운영진 {len(staff_members)}개, 대표자 {len(clubs)}개)')
if warnings:
    print(f'\n확인 필요 {len(warnings)}건:')
    for w in warnings:
        print(f'  - {w}')
